"""Generate synthetic educational fixtures on the CI host, never from school data."""
import base64
import gzip
from pathlib import Path
import subprocess
import sys
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches
from reportlab.pdfgen import canvas

root=Path(sys.argv[1]); root.mkdir(parents=True,exist_ok=True)
pdf=canvas.Canvas(str(root/'fractions.pdf')); pdf.setFont('Helvetica-Bold',18)
pdf.drawString(50,780,'Equivalent fractions'); pdf.setFont('Helvetica',12)
pdf.drawString(50,750,'One half is two quarters.'); pdf.save()
document=Document(); document.add_heading('Equivalent fractions',level=1)
document.add_paragraph('One half is two quarters.'); document.save(root/'fractions.docx')
slides=Presentation(); slide=slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text='Equivalent fractions'; slide.placeholders[1].text='One half is two quarters.'
slides.save(root/'fractions.pptx')
workbook=Workbook(); workbook.active.append(['Equivalent fractions','Whole'])
workbook.active.append(['One half',2]); workbook.save(root/'fractions.xlsx')
(root/'fractions.wav').write_bytes(gzip.decompress(base64.b64decode((Path(__file__).parent.parent/'tests/fixtures/fractions.wav.gz.b64').read_text())))
subprocess.run(['ffmpeg','-nostdin','-y','-v','error','-f','lavfi','-i','color=c=black:s=16x16:r=1','-i',str(root/'fractions.wav'),
    '-c:v','mpeg4','-c:a','alac','-shortest',str(root/'fractions.mp4')],check=True,timeout=30)
