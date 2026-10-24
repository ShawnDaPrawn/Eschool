import io
import subprocess
import time
import zipfile
import pytest
from navigateed_converter.archives import validate_file
from navigateed_converter.engine import Engine
from navigateed_converter.errors import ConversionError
from conftest import MemoryFetcher


def office_archive(extra_name=None, extra_data=b'x'):
    target = io.BytesIO()
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('[Content_Types].xml', '<Types/>')
        archive.writestr('word/document.xml', '<document/>')
        if extra_name: archive.writestr(extra_name, extra_data)
    return target.getvalue()

@pytest.mark.parametrize('name,data', [('../outside', b'x'), ('/outside', b'x'), ('word/../../outside', b'x'),
    ('word/vbaProject.bin', b'x'), ('word/embeddings/embedded.bin', b'x'), ('word/document2.xml', b'<!DOCTYPE x [<!ENTITY x SYSTEM "file:///etc/passwd">]><x/>'),
    ('word/document2.xml', '<!DOCTYPE x><x/>'.encode('utf-16')), ('word/bomb.xml', b'a'*300000)])
def test_unsafe_archives_are_refused_and_per_job_scratch_is_removed(settings, source, name, data):
    content = office_archive(name, data)
    with pytest.raises(ConversionError): Engine(settings, MemoryFetcher(content)).run(source(content, 'docx'))
    assert list(settings.scratch_root.iterdir()) == []


def test_actual_pdf_thirty_pages_is_readable_with_headings_and_within_five_minutes(settings, source):
    from reportlab.pdfgen.canvas import Canvas
    buffer = io.BytesIO()
    document = Canvas(buffer)
    for page in range(30):
        document.setFont('Helvetica-Bold', 22)
        document.drawString(50, 760, f'Fractions lesson {page+1}')
        document.setFont('Helvetica', 11)
        document.drawString(50, 720, 'One half and one quarter are examples of fractions.')
        document.showPage()
    document.save()
    content = buffer.getvalue()
    started = time.monotonic()
    result = Engine(settings, MemoryFetcher(content)).run(source(content, 'pdf'))
    assert time.monotonic()-started < 300
    assert '## Fractions lesson 1' in result['markdown']
    assert 'Fractions lesson 30' in result['markdown']
    assert 'One half and one quarter' in result['markdown']
    assert list(settings.scratch_root.iterdir()) == []


def test_actual_word_document_and_slides_keep_headings(settings, source):
    from docx import Document
    from pptx import Presentation
    word = Document()
    word.add_heading('Equivalent fractions', 1)
    word.add_paragraph('Two quarters make one half.')
    buffer = io.BytesIO(); word.save(buffer)
    content = buffer.getvalue()
    converted = Engine(settings, MemoryFetcher(content)).run(source(content, 'docx'))
    assert '# Equivalent fractions' in converted['markdown']
    assert 'Two quarters' in converted['markdown']
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = 'Fractions in diagrams'
    slide.placeholders[1].text = 'Shade one half of the diagram.'
    buffer = io.BytesIO(); slides.save(buffer)
    content = buffer.getvalue()
    converted = Engine(settings, MemoryFetcher(content)).run(source(content, 'pptx'))
    assert 'Fractions in diagrams' in converted['markdown']
    assert 'Shade one half' in converted['markdown']
    assert list(settings.scratch_root.iterdir()) == []


def test_failed_worker_and_timeout_remove_scratch_without_leaking_parser_errors(settings, source, monkeypatch):
    def timeout(*args, **kwargs): raise subprocess.TimeoutExpired('PRIVATE command', 1)
    monkeypatch.setattr('navigateed_converter.engine.subprocess.run', timeout)
    with pytest.raises(ConversionError, match='conversion_timeout'): Engine(settings, MemoryFetcher(b'lesson')).run(source(b'lesson'))
    assert list(settings.scratch_root.iterdir()) == []


def test_audio_cannot_bypass_gateway_and_missing_local_model_is_actionable(settings, source):
    content = b'RIFF' + b'0000' + b'WAVE'
    with pytest.raises(ConversionError, match='use_transcription_gateway'): Engine(settings).run(source(content, 'wav'))
    with pytest.raises(ConversionError, match='transcription_unavailable'): Engine(settings).run(source(content, 'wav'), True)
    assert list(settings.scratch_root.iterdir()) == []


def test_unnecessary_archive_codecs_are_refused_before_parent_process_expansion(tmp_path):
    target=tmp_path / 'lesson.docx'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_LZMA) as archive:
        archive.writestr('[Content_Types].xml','<Types/>')
        archive.writestr('word/document.xml','<document/>')
    with pytest.raises(ConversionError, match='unsafe_office_archive'):
        validate_file(target,'docx')
